from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from datetime import datetime, date
import re

BASE_DIR = Path(__file__).resolve().parents[2]
LOT_STATUS = BASE_DIR / "data" / "measurement" / "Lot_status.xlsx"


def _clean(v):
    return "" if v is None else str(v).strip()


def _parse_lot(v):
    m = re.search(r"(\d+)", _clean(v))
    return int(m.group(1)) if m else None


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = _clean(v)
    if not s:
        return None
    s = re.sub(r"(\d+)(st|nd|rd|th)", r"\1", s, flags=re.I)
    for fmt in ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except Exception:
            pass
    m = re.search(r"(\d{4})[-_/](\d{1,2})[-_/](\d{1,2})", s)
    if m:
        return f"{int(m.group(1)):04d}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    return None


def _parse_type(raw):
    s = _clean(raw)
    m = re.search(r"(\d+)\s*C", s, flags=re.I)
    count = int(m.group(1)) if m else None
    substrate = "base/unspecified"
    if re.search(r"SiO2", s, flags=re.I):
        substrate = "SiO2"
    elif re.search(r"(?:^|\W)Si(?:\W|$)", s, flags=re.I):
        substrate = "Si"
    return {
        "type_raw": s or None,
        "conditioning_count": count,
        "conditioning_suffix": substrate,
    }


@lru_cache(maxsize=1)
def load_lot_metadata():
    if not LOT_STATUS.exists():
        raise FileNotFoundError(f"Lot_status.xlsx not found: {LOT_STATUS}")

    from openpyxl import load_workbook
    wb = load_workbook(LOT_STATUS, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    header_idx = None
    headers = None
    for i, row in enumerate(rows[:20]):
        vals = [_clean(x) for x in row]
        if any("Lot No" in x for x in vals) and any(x == "Date" for x in vals):
            header_idx = i
            headers = vals
            break
    if header_idx is None:
        raise RuntimeError("Could not find Lot_status.xlsx header row.")

    normalized = {h.strip(): j for j, h in enumerate(headers) if h.strip()}

    def col(*names):
        for n in names:
            if n in normalized:
                return normalized[n]
        for h, j in normalized.items():
            for n in names:
                if n.lower() in h.lower():
                    return j
        return None

    lot_i = col("Lot No.", "Lot No", "Lot")
    date_i = col("Date")
    waf_i = col("Wafers", "Wafers ")
    meas_i = col("Measurements")
    type_i = col("Type")

    records = []
    for row in rows[header_idx + 1:]:
        lot = _parse_lot(row[lot_i] if lot_i is not None and lot_i < len(row) else None)
        dt = _parse_date(row[date_i] if date_i is not None and date_i < len(row) else None)
        if lot is None or dt is None:
            continue
        expected = None
        if waf_i is not None and waf_i < len(row):
            try:
                expected = int(float(row[waf_i]))
            except Exception:
                expected = None
        measurements = _clean(row[meas_i]) if meas_i is not None and meas_i < len(row) else None
        type_raw = row[type_i] if type_i is not None and type_i < len(row) else None
        info = _parse_type(type_raw)
        records.append({
            "lot": lot,
            "date": dt,
            "expected_wafers": expected,
            "measurements": measurements or None,
            **info,
        })

    records.sort(key=lambda x: x["lot"])
    by_date = {r["date"]: r for r in records}
    by_lot = {r["lot"]: r for r in records}
    return {"records": records, "by_date": by_date, "by_lot": by_lot}


def experiment_context(experiment: str):
    m = re.match(r"(\d{4}-\d{2}-\d{2})_(\d+)$", experiment)
    if not m:
        raise KeyError(experiment)
    dt, wafer = m.group(1), int(m.group(2))
    meta = load_lot_metadata()
    row = meta["by_date"].get(dt)
    if row is None:
        raise KeyError(f"No Lot_status metadata for {dt}")
    return {
        "experiment": experiment,
        "wafer": wafer,
        **row,
        "cohort": "primary_1_9" if row["lot"] <= 9 else "external_lot10",
    }


def all_metadata():
    return load_lot_metadata()["records"]
